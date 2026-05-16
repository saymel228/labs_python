import tkinter as tk
from tkinter import ttk, messagebox
import parallelepiped
import tetrahedron
import sphere
from docx import Document
from openpyxl import Workbook

class GeometryCalculator:
    def __init__(self, root):
        self.root = root
        self.root.title("Калькулятор геометрических тел")
        self.root.geometry("500x400")
        
        self.figure_var = tk.StringVar()
        self.figure_var.set("Параллелепипед")
        
        figures = ["Параллелепипед", "Тетраэдр", "Шар"]
        ttk.Label(root, text="Выберите фигуру:").pack(pady=5)
        self.figure_menu = ttk.Combobox(root, textvariable=self.figure_var, values=figures)
        self.figure_menu.pack(pady=5)
        self.figure_menu.bind("<<ComboboxSelected>>", self.update_parameters)
        
        self.parameters_frame = ttk.Frame(root)
        self.parameters_frame.pack(pady=10)
        self.parameter_entries = []
        
        self.result_text = tk.Text(root, height=10, width=50)
        self.result_text.pack(pady=10)
        
        ttk.Button(root, text="Рассчитать", command=self.calculate).pack(pady=5)
        ttk.Button(root, text="Сохранить в DOCX", command=self.save_to_docx).pack(side=tk.LEFT, padx=20)
        ttk.Button(root, text="Сохранить в XLSX", command=self.save_to_xlsx).pack(side=tk.RIGHT, padx=20)
        
        self.update_parameters()

    def update_parameters(self, event=None):
        for widget in self.parameters_frame.winfo_children():
            widget.destroy()
        self.parameter_entries = []
        
        figure = self.figure_var.get()
        
        if figure == "Параллелепипед":
            params = parallelepiped.get_parameters()
            self.volume_func = parallelepiped.calculate_volume
            self.area_func = parallelepiped.calculate_surface_area
        elif figure == "Тетраэдр":
            params = tetrahedron.get_parameters()
            self.volume_func = tetrahedron.calculate_volume
            self.area_func = tetrahedron.calculate_surface_area
        elif figure == "Шар":
            params = sphere.get_parameters()
            self.volume_func = sphere.calculate_volume
            self.area_func = sphere.calculate_surface_area
        
        for param in params:
            frame = ttk.Frame(self.parameters_frame)
            frame.pack(fill=tk.X, pady=2)
            
            ttk.Label(frame, text=param, width=15).pack(side=tk.LEFT)
            entry = ttk.Entry(frame)
            entry.pack(side=tk.RIGHT, expand=True, fill=tk.X)
            self.parameter_entries.append(entry)

    def calculate(self):
        try:
            values = [float(entry.get()) for entry in self.parameter_entries]
            figure = self.figure_var.get()
            
            volume = self.volume_func(*values)
            area = self.area_func(*values)
            
            result = f"Фигура: {figure}\n"
            params = parallelepiped.get_parameters() if figure == "Параллелепипед" else \
                     tetrahedron.get_parameters() if figure == "Тетраэдр" else \
                     sphere.get_parameters()
            
            for i, param in enumerate(params):
                result += f"{param}: {values[i]}\n"
            result += f"\nОбъём: {volume:.2f}\n"
            result += f"Площадь поверхности: {area:.2f}"
            
            self.result_text.delete(1.0, tk.END)
            self.result_text.insert(tk.END, result)
            
        except ValueError:
            messagebox.showerror("Ошибка", "Пожалуйста, введите корректные числовые значения")

    def save_to_docx(self):
        doc = Document()
        doc.add_heading('Результаты расчёта', 0)
        doc.add_paragraph(self.result_text.get(1.0, tk.END))
        doc.save('geometry_report.docx')
        messagebox.showinfo("Успех", "Файл сохранён как geometry_report.docx")

    def save_to_xlsx(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Результаты"
        
        for i, line in enumerate(self.result_text.get(1.0, tk.END).split('\n')):
            ws.cell(row=i+1, column=1, value=line)
        
        wb.save('geometry_report.xlsx')
        messagebox.showinfo("Успех", "Файл сохранён как geometry_report.xlsx")

if __name__ == "__main__":
    root = tk.Tk()
    app = GeometryCalculator(root)
    root.mainloop()