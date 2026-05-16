from kivy.app import App
from kivy.uix.boxlayout import BoxLayout
from kivy.uix.label import Label
from kivy.uix.textinput import TextInput
from kivy.uix.spinner import Spinner
from kivy.uix.button import Button
from kivy.uix.popup import Popup
from kivy.core.window import Window
from abc import ABC, abstractmethod
import math
from docx import Document
from openpyxl import Workbook


class GeometryBody(ABC):
    @property
    @abstractmethod
    def parameters(self) -> list:
        pass
    
    @abstractmethod
    def volume(self) -> float:
        pass
    
    @abstractmethod
    def surface_area(self) -> float:
        pass
    
    def __str__(self):
        return f"{self.__class__.__name__}: объём = {self.volume():.2f}, площадь = {self.surface_area():.2f}"
    
    def __repr__(self):
        return f"{self.__class__.__name__}({self._get_params_repr()})"
    
    @abstractmethod
    def _get_params_repr(self) -> str:
        pass
    
    def __eq__(self, other):
        if not isinstance(other, GeometryBody):
            return False
        return self.volume() == other.volume()
    
    def __lt__(self, other):
        if not isinstance(other, GeometryBody):
            raise TypeError("Сравнение возможно только с геометрическими телами")
        return self.volume() < other.volume()


class Parallelepiped(GeometryBody):
    def __init__(self, length=1.0, width=1.0, height=1.0):
        self.length = length
        self.width = width
        self.height = height
    
    @property
    def parameters(self) -> list:
        return ["Длина", "Ширина", "Высота"]
    
    @property
    def length(self):
        return self._length
    
    @length.setter
    def length(self, value):
        if value <= 0:
            raise ValueError("Длина должна быть положительным числом")
        self._length = float(value)
    
    @property
    def width(self):
        return self._width
    
    @width.setter
    def width(self, value):
        if value <= 0:
            raise ValueError("Ширина должна быть положительным числом")
        self._width = float(value)
    
    @property
    def height(self):
        return self._height
    
    @height.setter
    def height(self, value):
        if value <= 0:
            raise ValueError("Высота должна быть положительным числом")
        self._height = float(value)
    
    def volume(self) -> float:
        return self.length * self.width * self.height
    
    def surface_area(self) -> float:
        return 2 * (self.length*self.width + self.length*self.height + self.width*self.height)
    
    def _get_params_repr(self) -> str:
        return f"length={self.length}, width={self.width}, height={self.height}"
    
    def __add__(self, other):
        if not isinstance(other, Parallelepiped):
            raise TypeError("Можно складывать только параллелепипеды")
        return self.volume() + other.volume()
    
    def __mul__(self, factor):
        if not isinstance(factor, (int, float)):
            raise TypeError("Коэффициент должен быть числом")
        return Parallelepiped(self.length * factor, self.width * factor, self.height * factor)


class Tetrahedron(GeometryBody):
    def __init__(self, edge=1.0):
        self.edge = edge
    
    @property
    def parameters(self) -> list:
        return ["Длина ребра"]
    
    @property
    def edge(self):
        return self._edge
    
    @edge.setter
    def edge(self, value):
        if value <= 0:
            raise ValueError("Длина ребра должна быть положительным числом")
        self._edge = float(value)
    
    def volume(self) -> float:
        return (self.edge ** 3) * math.sqrt(2) / 12
    
    def surface_area(self) -> float:
        return math.sqrt(3) * (self.edge ** 2)
    
    def _get_params_repr(self) -> str:
        return f"edge={self.edge}"
    
    def __sub__(self, other):
        if not isinstance(other, Tetrahedron):
            raise TypeError("Можно вычитать только тетраэдры")
        result = self.volume() - other.volume()
        return max(result, 0)
    
    def __pow__(self, power):
        if not isinstance(power, (int, float)):
            raise TypeError("Степень должна быть числом")
        return Tetrahedron(self.edge ** power)


class Sphere(GeometryBody):
    def __init__(self, radius=1.0):
        self.radius = radius
    
    @property
    def parameters(self) -> list:
        return ["Радиус"]
    
    @property
    def radius(self):
        return self._radius
    
    @radius.setter
    def radius(self, value):
        if value <= 0:
            raise ValueError("Радиус должен быть положительным числом")
        self._radius = float(value)
    
    def volume(self) -> float:
        return (4/3) * math.pi * (self.radius ** 3)
    
    def surface_area(self) -> float:
        return 4 * math.pi * (self.radius ** 2)
    
    def _get_params_repr(self) -> str:
        return f"radius={self.radius}"
    
    def __truediv__(self, divisor):
        if not isinstance(divisor, (int, float)):
            raise TypeError("Делитель должен быть числом")
        if divisor == 0:
            raise ValueError("Деление на ноль невозможно")
        return Sphere(self.radius / divisor)
    
    def __contains__(self, point):
        if not isinstance(point, (tuple, list)) or len(point) != 3:
            raise TypeError("Точка должна быть кортежем (x, y, z)")
        x, y, z = point
        distance = math.sqrt(x**2 + y**2 + z**2)
        return distance <= self.radius


class GeometryApp(App):
    def build(self):
        Window.size = (550, 700)
        
        main_layout = BoxLayout(orientation='vertical', padding=20, spacing=15)
        
        main_layout.add_widget(Label(text="Выберите геометрическое тело:", font_size=20, size_hint_y=None, height=40))
        
        self.spinner = Spinner(
            values=["Параллелепипед", "Тетраэдр", "Сфера"],
            size_hint=(1, None),
            height=50,
            font_size=18,
            text="Параллелепипед"
        )
        main_layout.add_widget(self.spinner)
        
        self.params_layout = BoxLayout(orientation='vertical', spacing=15, size_hint_y=None)
        self.params_layout.bind(minimum_height=self.params_layout.setter('height'))
        main_layout.add_widget(self.params_layout)

        btn_calc = Button(
            text="Рассчитать", 
            size_hint=(1, None), 
            height=60,
            font_size=18
        )
        btn_calc.bind(on_press=self.calculate)
        main_layout.add_widget(btn_calc)

        btn_box = BoxLayout(size_hint=(1, None), height=70, spacing=15)
        btn_docx = Button(text="Экспорт в DOCX", font_size=16)
        btn_docx.bind(on_press=lambda x: self.export('docx'))
        btn_xlsx = Button(text="Экспорт в XLSX", font_size=16)
        btn_xlsx.bind(on_press=lambda x: self.export('xlsx'))
        btn_box.add_widget(btn_docx)
        btn_box.add_widget(btn_xlsx)
        main_layout.add_widget(btn_box)
        
        self.result_label = Label(
            size_hint_y=None,
            height=220,
            font_size=16,
            halign='left',
            valign='top'
        )
        self.result_label.bind(size=self.result_label.setter('text_size'))
        main_layout.add_widget(self.result_label)

        self.spinner.bind(text=self.update_parameters)
     
        self.update_parameters(None, self.spinner.text)
        
        return main_layout
    
    def update_parameters(self, instance, value):
        """Обновляет поля ввода параметров при смене фигуры"""
        if not value or value == '':
            return
        
        self.params_layout.clear_widgets()
        
        body_classes = {
            "Параллелепипед": Parallelepiped,
            "Тетраэдр": Tetrahedron,
            "Сфера": Sphere
        }

        if value not in body_classes:
            return
            
        body_class = body_classes[value]
        params = body_class().parameters

        for param in params:
            param_box = BoxLayout(orientation='horizontal', size_hint_y=None, height=60)
            
            param_label = Label(
                text=param + ":", 
                size_hint_x=0.4,
                font_size=18,
                halign='right',
                valign='middle'
            )
            param_label.bind(size=param_label.setter('text_size'))
            param_box.add_widget(param_label)

            text_input = TextInput(
                multiline=False,
                font_size=18,
                size_hint_x=0.6,
                height=50,
                text="1.0"
            )
            param_box.add_widget(text_input)
            
            self.params_layout.add_widget(param_box)
    
    def calculate(self, instance):
        try:
            body_classes = {
                "Параллелепипед": Parallelepiped,
                "Тетраэдр": Tetrahedron,
                "Сфера": Sphere
            }
            
            body_class = body_classes[self.spinner.text]

            values = []
            for child in reversed(self.params_layout.children):
                if isinstance(child, BoxLayout):
                    for widget in child.children:
                        if isinstance(widget, TextInput):
                            values.append(float(widget.text))
                            break
            
            self.current_body = body_class(*values)
            
            result = f"Тип: {self.spinner.text}\n\n"
            result += "Параметры:\n"
            
            params = body_class().parameters
            for i, param in enumerate(params):
                result += f"  {param}: {values[i]}\n"
            
            result += f"\nОбъём: {self.current_body.volume():.2f}\n"
            result += f"Площадь поверхности: {self.current_body.surface_area():.2f}\n\n"
            result += f"{self.current_body}"
            
            self.result_label.text = result
            
        except ValueError as e:
            self.show_error(f"Ошибка ввода: {str(e)}")
        except Exception as e:
            self.show_error(f"Ошибка: {str(e)}")
    
    def export(self, format_type):
        if not hasattr(self, 'current_body') or self.current_body is None:
            self.show_error("Сначала выполните расчет")
            return
        
        try:
            body = self.current_body

            if isinstance(body, Parallelepiped):
                params_dict = {
                    "Длина": body.length,
                    "Ширина": body.width,
                    "Высота": body.height
                }
            elif isinstance(body, Tetrahedron):
                params_dict = {
                    "Длина ребра": body.edge
                }
            elif isinstance(body, Sphere):
                params_dict = {
                    "Радиус": body.radius
                }
            else:
                params_dict = {}
            
            data = {
                "type": self.spinner.text,
                "params": params_dict,
                "volume": body.volume(),
                "area": body.surface_area()
            }
            
            if format_type == 'docx':
                doc = Document()
                doc.add_heading('Результаты расчета', 0)
                doc.add_paragraph(f"Тип тела: {data['type']}")
                doc.add_paragraph("Параметры:")
                for param, value in data['params'].items():
                    doc.add_paragraph(f"{param}: {value}")
                doc.add_paragraph(f"Объём: {data['volume']:.2f}")
                doc.add_paragraph(f"Площадь поверхности: {data['area']:.2f}")
                doc.add_paragraph(f"Представление: {body}")
                doc.save('geometry_report.docx')
                self.show_message("Результаты сохранены в geometry_report.docx")
            else:
                wb = Workbook()
                ws = wb.active
                ws.title = "Результаты"
                ws.append(["Тип тела", data['type']])
                ws.append(["Параметры:"])
                for param, value in data['params'].items():
                    ws.append([param, value])
                ws.append(["Объём", f"{data['volume']:.2f}"])
                ws.append(["Площадь поверхности", f"{data['area']:.2f}"])
                ws.append(["Представление", str(body)])
                wb.save('geometry_report.xlsx')
                self.show_message("Результаты сохранены в geometry_report.xlsx")
                
        except Exception as e:
            self.show_error(f"Ошибка экспорта: {str(e)}")
    
    def show_error(self, message):
        popup = Popup(
            title="Ошибка",
            content=Label(text=message, font_size=18),
            size_hint=(0.7, 0.4)
        )
        popup.open()
    
    def show_message(self, message):
        popup = Popup(
            title="Успешно",
            content=Label(text=message, font_size=18),
            size_hint=(0.7, 0.4)
        )
        popup.open()


if __name__ == '__main__':
    GeometryApp().run()